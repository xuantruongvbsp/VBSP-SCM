"""Regression tests for CBTD dashboard alert summary/export helpers."""
from __future__ import annotations

from io import BytesIO

import pandas as pd
from openpyxl import load_workbook

import tabs.tab_cbtd_dashboard as dashboard


def _alerts() -> list[dict]:
    return [
        {
            "loai": "cbtd_qh_cao",
            "muc_do": "🔴",
            "noi_dung": "CBTD **CB01 — A** có tỷ lệ QH 3.0%",
            "chi_tiet": {
                "ma_cb": "CB01",
                "ho_ten": "A",
                "ty_le_qh": 3.0,
                "tong_du_no": 100_000_000,
                "du_no_qh": 3_000_000,
            },
        },
        {
            "loai": "to_yeu_lien_tiep",
            "muc_do": "🔴",
            "noi_dung": "Tổ **T01** xếp loại Yếu",
            "chi_tiet": {
                "ten_dv": "PGD A",
                "ten_xa": "Xã A",
                "ma_to": "T01",
                "xep_loai": "Yếu",
                "tong_diem": 35,
            },
        },
        {
            "loai": "dgd_thieu_cbtd",
            "muc_do": "⚠️",
            "noi_dung": "ĐGD **D1** chưa có CBTD phụ trách",
            "chi_tiet": {"pgd": "PGD B", "xa": "Xã B", "dgd": "D1"},
        },
        {
            "loai": r"loai/la?[]:*\\rat_dai_de_test_ten_sheet_excel",
            "muc_do": "⚠️",
            "noi_dung": pd.NA,
            "chi_tiet": {"pgd": pd.NA, "foo": "bar"},
        },
    ]


def test_bang_tong_hop_canh_bao_fallback_don_vi_va_bo_nan_gia():
    df = dashboard._bang_tong_hop_canh_bao(_alerts(), {"CB01": {"pgd": "PGD A"}})

    by_type = df.set_index("Loại cảnh báo")
    assert by_type.loc["CBTD QH cao", "Đơn vị nhiều nhất"] == "PGD A (1)"
    assert by_type.loc["CBTD QH cao", "Số đơn vị"] == 1
    assert by_type.loc["ĐGD thiếu CBTD", "Đơn vị nhiều nhất"] == "PGD B (1)"
    assert not df.astype(str).apply(lambda col: col.str.contains(r"\bnan\b|<NA>", case=False)).any().any()
    assert df.iloc[0]["Mức độ"] == "🔴"


def test_xlsx_chi_tiet_canh_bao_sheet_hop_le_va_noi_dung_sach():
    data = dashboard._xlsx_chi_tiet_canh_bao(_alerts(), {"CB01": {"pgd": "PGD A"}})
    wb = load_workbook(BytesIO(data), read_only=True, data_only=True)

    assert "Tổng hợp" in wb.sheetnames
    assert "Tất cả cảnh báo" in wb.sheetnames
    assert "CBTD QH cao" in wb.sheetnames
    assert "Tổ TB_Yếu 2+ kỳ" in wb.sheetnames
    assert all(len(name) <= 31 and not set(name) & dashboard._EXCEL_SHEET_CHARS_CAM for name in wb.sheetnames)

    ws_qh = wb["CBTD QH cao"]
    rows_qh = list(ws_qh.iter_rows(values_only=True))
    assert rows_qh[0][:3] == ("Đơn vị", "Mã CBTD", "Họ tên")
    assert rows_qh[1][0] == "PGD A"
    assert "**" not in str(rows_qh[1][-1])

    ws_all = wb["Tất cả cảnh báo"]
    rows_all = list(ws_all.iter_rows(values_only=True))
    assert rows_all[1][2] == "PGD A"
    assert "**" not in str(rows_all[1][3])
    assert all("<NA>" not in str(cell) and str(cell).lower() != "nan" for row in rows_all for cell in row)


def test_doc_cdtotkvv_moi_nhat_uu_tien_nguon_trung_tam(monkeypatch):
    from services import cdtotkvv_service, tongquan_cdto_service

    df_trung_tam = pd.DataFrame({"ten_dv": ["PGD A"], "ten_xa": ["Xã A"], "ma_to": ["T01"]})
    df_fallback = pd.DataFrame({"ten_dv": ["PGD B"], "ten_xa": ["Xã B"], "ma_to": ["T02"]})
    monkeypatch.setattr(tongquan_cdto_service, "load_cdto_toan_cn", lambda: {"df_raw": df_trung_tam})
    monkeypatch.setattr(cdtotkvv_service, "tong_hop_tu_pgd_data", lambda: df_fallback)

    result = dashboard._doc_cdtotkvv_moi_nhat()

    assert result is df_trung_tam


def test_doc_cdtotkvv_moi_nhat_fallback_khi_nguon_trung_tam_rong(monkeypatch):
    from services import cdtotkvv_service, tongquan_cdto_service

    df_fallback = pd.DataFrame({"ten_dv": ["PGD B"], "ten_xa": ["Xã B"], "ma_to": ["T02"]})
    monkeypatch.setattr(tongquan_cdto_service, "load_cdto_toan_cn", lambda: {"df_raw": pd.DataFrame()})
    monkeypatch.setattr(cdtotkvv_service, "tong_hop_tu_pgd_data", lambda: df_fallback)

    result = dashboard._doc_cdtotkvv_moi_nhat()

    assert result is df_fallback


def test_loc_df_cdto_khop_pgd_bang_ma_dv_va_xa_co_tien_to():
    df_cdto = pd.DataFrame({
        "ma_dv": ["004602", "004602", "004603"],
        "ten_dv": ["Tên lệch", "Tên lệch", "PGD Trảng Bom"],
        "ten_xa": ["Xã Long Đức", "Xã Khác", "Xã Long Đức"],
        "ma_to": ["T01", "T02", "T03"],
    })
    cbtd_data = {"CB01": {"pgd": "PGD Long Thành", "ds_dgd": ["ĐGD 1"]}}
    dgd_map = {"PGD Long Thành": {"Long Đức": {"ĐGD 1": {}}}}

    result = dashboard._loc_df_cdto(df_cdto, cbtd_data, dgd_map, "PGD Long Thành", "CB01")

    assert result["ma_to"].tolist() == ["T01"]


def test_loc_df_cdto_chon_cbtd_khoa_theo_pgd_truoc_khi_loc_xa():
    df_cdto = pd.DataFrame({
        "ma_dv": ["004602", "004603"],
        "ten_dv": ["Tên lệch A", "PGD Trảng Bom"],
        "ten_xa": ["Xã Long Đức", "Xã Long Đức"],
        "ma_to": ["T01", "T02"],
    })
    cbtd_data = {
        "CB01": {"pgd": "PGD Long Thành", "ds_dgd": ["ĐGD 1"]},
    }
    dgd_map = {
        "PGD Long Thành": {
            "Long Đức": {" ĐGD 1 ": {}},
        },
        "PGD Trảng Bom": {
            "Long Đức": {"ĐGD khác": {}},
        },
    }

    result = dashboard._loc_df_cdto(df_cdto, cbtd_data, dgd_map, "(Tất cả)", "CB01")

    assert result["ma_to"].tolist() == ["T01"]
